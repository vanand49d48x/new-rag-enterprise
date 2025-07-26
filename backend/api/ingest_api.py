# 5. ✅ FastAPI Endpoint to Trigger Ingestion
# File: backend/api/ingest_api.py

from fastapi import APIRouter, UploadFile, File
from backend.ingest.processor import process_document
from backend.rag.vector_store import index_chunks
import mimetypes
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
import csv
import io
import openpyxl
import email
from bs4 import BeautifulSoup
from typing import List

router = APIRouter()

@router.post("/ingest")
def ingest_files(files: List[UploadFile] = File(...)):
    total_chunks = 0
    results = []

    for file in files:
        filename = file.filename.lower()

        if filename.endswith(".pdf"):
            reader = PdfReader(file.file)
            content = "\n\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
        elif filename.endswith(".docx") or filename.endswith(".doc"):
            doc = DocxDocument(file.file)
            content = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        elif filename.endswith(".csv"):
            decoded = file.file.read().decode("utf-8")
            csv_reader = csv.reader(io.StringIO(decoded))
            content = "\n".join([", ".join(row) for row in csv_reader])
        elif filename.endswith(".xlsx"):
            workbook = openpyxl.load_workbook(file.file, read_only=True)
            sheet_text = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    sheet_text.append(", ".join([str(cell) if cell is not None else "" for cell in row]))
            content = "\n".join(sheet_text)
        elif filename.endswith(".html") or filename.endswith(".htm"):
            html = file.file.read().decode("utf-8")
            soup = BeautifulSoup(html, "html.parser")
            content = soup.get_text(separator="\n")
        elif filename.endswith(".rtf"):
            from striprtf.striprtf import rtf_to_text
            rtf = file.file.read().decode("utf-8")
            content = rtf_to_text(rtf)
        elif filename.endswith(".eml"):
            msg = email.message_from_binary_file(file.file)
            content = ""
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        content += part.get_payload(decode=True).decode("utf-8")
            else:
                content = msg.get_payload(decode=True).decode("utf-8")
        else:
            content = file.file.read().decode("utf-8")

        chunks = process_document(
            text=content,
            title=file.filename,
            doc_type="uploaded"
        )
        index_chunks(chunks)
        total_chunks += len(chunks)
        results.append({"file": file.filename, "chunks": len(chunks)})

    return {"message": f"{total_chunks} total chunks ingested", "details": results}
